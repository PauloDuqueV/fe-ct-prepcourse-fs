"""Genera Inventario_Real_Meditienda.xlsx a partir de los informes Siigo (uno por mes)."""
import sys, os, re, glob
from collections import Counter
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from procesar_inventario import cargar, analizar, NO_FISICOS_PREFIJO, NO_FISICOS_CODIGO

ARIAL = "Arial"
GRIS = PatternFill("solid", fgColor="44546A")
TEAL = PatternFill("solid", fgColor="0097A7")
GRIS2 = PatternFill("solid", fgColor="D9E1E8")
AMAR = PatternFill("solid", fgColor="FFF2CC")
ROJO_F = PatternFill("solid", fgColor="F8D7DA")
fino = Side(style="thin", color="7F7F7F")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
IZQ = Alignment(horizontal="left", vertical="center", wrap_text=True)
F = lambda **k: Font(name=ARIAL, size=k.pop("size", 10), **k)

def encabezado(ws, fila, textos, relleno=GRIS):
    for i, t in enumerate(textos, 1):
        c = ws.cell(fila, i, t); c.font = F(bold=True, color="FFFFFF"); c.fill = relleno
        c.alignment = CENTRO; c.border = BORDE

def anchos(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = x

def leer_maestro(path):
    if not path: return pd.Series(dtype=str)
    c = pd.read_excel(path, dtype=str).iloc[:, :2]
    c.columns = ["cod", "nom"]
    c["cod"] = c["cod"].str.strip().str.upper()
    c["nom"] = c["nom"].fillna("").map(lambda x: " ".join(x.split()))
    return c.drop_duplicates("cod").set_index("cod")["nom"]

def main(files, salida, maestro_path=None):
    d, total_filas = cargar(files)
    d, L, R, nombres = analizar(d)
    maestro = leer_maestro(maestro_path)
    R["En maestro"] = R.index.isin(maestro.index)
    sin_mov = [c for c in maestro.index if c not in R.index
               and not (c.startswith(NO_FISICOS_PREFIJO) or c in NO_FISICOS_CODIGO)]
    meses = list(dict.fromkeys(d.sort_values("Fecha")["Mes"]))
    fis = R[R["Físico"]]
    no_fis = R[~R["Físico"]]
    rank = R["Ranking"].to_dict()
    L["Ranking"] = L["Código"].map(rank)
    L = L.sort_values(["Ranking", "Orden"])

    wb = Workbook()
    # ======================= INVENTARIO =======================
    ws = wb.active; ws.title = "Inventario"
    ws["A1"] = "INVENTARIO REAL – DISTRIBUCIONES MEDITIENDA S.A.S"; ws["A1"].font = F(bold=True, size=14)
    ws["A2"] = (f"Base: informes Siigo de {meses[0].title()} a {meses[-1].title()} 2026. Productos ordenados de mayor a menor rotación. "
                "Al final van los productos del maestro sin movimiento en el periodo. Llene solo las celdas amarillas (Cantidad por lote) con el conteo físico de bodega. Detalle en la hoja 'Resumen'.")
    ws["A2"].font = F(italic=True, size=9); ws.merge_cells("A2:J2"); ws["A2"].alignment = IZQ; ws.row_dimensions[2].height = 28
    ws.merge_cells("A3:D3"); ws["A3"] = "PRODUCTO"
    ws.merge_cells("E3:H3"); ws["E3"] = "SEGUIMIENTO LOTE"
    ws.merge_cells("I3:J3"); ws["I3"] = "CONTEO"
    for col, fill in (("A", GRIS), ("E", TEAL), ("I", GRIS)):
        c = ws[f"{col}3"]; c.font = F(bold=True, color="FFFFFF", size=11); c.alignment = CENTRO
    for c in range(1, 11):
        ws.cell(3, c).fill = GRIS if c <= 4 or c >= 9 else TEAL; ws.cell(3, c).border = BORDE
    hdr = ["Prioridad", "Código", "Nombre del producto", "Opción", "Lote", "Fecha vencimiento", "Registro Invima",
           "Cantidad por lote", "Total contado", "Observación"]
    encabezado(ws, 4, hdr)
    for c in range(5, 9): ws.cell(4, c).fill = TEAL
    sel = L[L["Orden"] <= 2].set_index(["Código", "Orden"])
    r = 5
    filas_inv = [(cod, int(row["Ranking"]), bool(row["En maestro"]) or maestro.empty) for cod, row in fis.iterrows()]
    filas_inv += [(cod, len(R) + i, True) for i, cod in enumerate(sin_mov, 1)]
    for cod, ranking, en_maestro in filas_inv:
        r1, r2 = r, r + 1
        for col in ("A", "B", "C", "I"): ws.merge_cells(f"{col}{r1}:{col}{r2}")
        ws[f"A{r1}"] = ranking; ws[f"B{r1}"] = cod
        ws[f"C{r1}"] = f'=IF(IFERROR(INDEX(Catalogo!$B:$B,MATCH(B{r1},Catalogo!$A:$A,0)),"")="","(nombre pendiente – ver Catalogo)",INDEX(Catalogo!$B:$B,MATCH(B{r1},Catalogo!$A:$A,0)))'
        ws[f"I{r1}"] = f'=IF(COUNT(H{r1}:H{r2})=0,"",SUM(H{r1}:H{r2}))'
        for k, rr in ((1, r1), (2, r2)):
            ws[f"D{rr}"] = k
            obs = []
            if (cod, k) in sel.index:
                s = sel.loc[(cod, k)]
                ws[f"E{rr}"] = s["Lote"]; ws[f"F{rr}"] = s["FV"]; ws[f"G{rr}"] = s["Invima"]
                if s["Vencido"]: obs.append("Lote vencido")
                if pd.isna(s["Última compra"]): obs.append("Lote visto solo en ventas")
                if s["Notas"] and ("distint" in s["Notas"] or "inválida" in s["Notas"] or "varios lotes" in s["Notas"]):
                    obs.append("Verificar (ver hoja Lotes)")
            elif k == 1:
                obs.append("Sin movimiento jun–oct" if cod in sin_mov else "Sin lote registrado en los informes")
            if k == 1 and not en_maestro: obs.insert(0, "Código no está en el maestro")
            ws[f"J{rr}"] = "; ".join(obs)
            ws[f"H{rr}"].fill = AMAR
        for rr in (r1, r2):
            for c in range(1, 11):
                cell = ws.cell(rr, c); cell.border = BORDE; cell.font = F()
                cell.alignment = IZQ if c in (3, 10) else CENTRO
            ws.row_dimensions[rr].height = 18
        ws[f"I{r1}"].font = F(bold=True)
        r += 2
    ultima = r - 1
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        error="Escriba una cantidad mayor o igual a 0", errorTitle="Cantidad no válida")
    ws.add_data_validation(dv); dv.add(f"H5:H{ultima}")
    ws.conditional_formatting.add(f"E5:G{ultima}", FormulaRule(formula=[f'ISNUMBER(SEARCH("vencido",$J5))'], fill=ROJO_F))
    anchos(ws, [9, 11, 52, 8, 18, 12, 22, 12, 11, 30])
    ws.freeze_panes = "C5"; ws.auto_filter.ref = f"A4:J{ultima}"
    ws.print_title_rows = "3:4"; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ======================= ROTACION =======================
    wr = wb.create_sheet("Rotacion")
    wr["A1"] = "Rotación por código (entradas = Factura de compra, salidas = Factura de venta; columna V 'Cantidad')"
    wr["A1"].font = F(bold=True, size=12)
    encabezado(wr, 3, ["Ranking", "Código", "Nombre del producto", "Entradas (unid.)", "Salidas (unid.)",
                       "Neto entradas – salidas", "N.º compras", "N.º ventas", "Total movido (unid.)", "Lotes distintos", "En plantilla", "En maestro"])
    n_lotes = L.groupby("Código").size()
    for i, (cod, row) in enumerate(R.iterrows(), start=4):
        wr.cell(i, 1, int(row["Ranking"])); wr.cell(i, 2, cod)
        wr.cell(i, 3, f'=IFERROR(INDEX(Catalogo!$B:$B,MATCH(B{i},Catalogo!$A:$A,0))&"","")')
        wr.cell(i, 4, f'=SUMIFS(Datos!$H:$H,Datos!$F:$F,B{i},Datos!$B:$B,"Factura de compra")')
        wr.cell(i, 5, f'=SUMIFS(Datos!$H:$H,Datos!$F:$F,B{i},Datos!$B:$B,"Factura de venta")')
        wr.cell(i, 6, f"=D{i}-E{i}")
        wr.cell(i, 7, f'=COUNTIFS(Datos!$F:$F,B{i},Datos!$B:$B,"Factura de compra")')
        wr.cell(i, 8, f'=COUNTIFS(Datos!$F:$F,B{i},Datos!$B:$B,"Factura de venta")')
        wr.cell(i, 9, f"=D{i}+E{i}")
        wr.cell(i, 10, int(n_lotes.get(cod, 0)))
        wr.cell(i, 11, "Sí" if row["Físico"] else "No (servicio/flete)")
        wr.cell(i, 12, "Sí" if row["En maestro"] else "No")
        for c in range(1, 13):
            cell = wr.cell(i, c); cell.font = F(); cell.border = BORDE
            if c in (4, 5, 6, 9): cell.number_format = "#,##0.##;(#,##0.##);-"
    anchos(wr, [9, 11, 50, 14, 14, 16, 11, 11, 16, 11, 18, 11])
    wr.freeze_panes = "C4"; wr.auto_filter.ref = f"A3:L{3 + len(R)}"

    # ======================= LOTES =======================
    wl = wb.create_sheet("Lotes")
    wl["A1"] = "Lotes homologados por producto (todas las escrituras de un mismo lote se unifican en una sola fila)"
    wl["A1"].font = F(bold=True, size=12)
    cols = ["Ranking", "Código", "Lote", "FV", "Invima", "Última compra", "Última venta", "Unid. compradas",
            "Unid. vendidas", "Registros", "Escrituras distintas", "Selección", "Criterio de orden", "Vencido", "Notas de homologación", "Ejemplos de escritura (col. Q)"]
    encabezado(wl, 3, cols)
    for i, (_, x) in enumerate(L.iterrows(), start=4):
        vals = [int(x["Ranking"]), x["Código"], x["Lote"], x["FV"], x["Invima"],
                None if pd.isna(x["Última compra"]) else x["Última compra"].to_pydatetime(),
                None if pd.isna(x["Última venta"]) else x["Última venta"].to_pydatetime(),
                x["Unid. compradas"], x["Unid. vendidas"], int(x["Registros"]), int(x["Escrituras distintas"]),
                x["Seleccionado"] or "No (más antiguo)", x["Criterio"], "Sí" if x["Vencido"] else "", x["Notas"], x["Ejemplos de escritura"]]
        for c, v in enumerate(vals, 1):
            cell = wl.cell(i, c, v); cell.font = F(); cell.border = BORDE
            if c in (6, 7): cell.number_format = "DD/MM/YYYY"
        if x["Seleccionado"]:
            for c in range(1, 13): wl.cell(i, c).fill = GRIS2
    anchos(wl, [9, 11, 18, 10, 22, 12, 12, 11, 11, 9, 11, 16, 22, 8, 55, 70])
    wl.freeze_panes = "D4"; wl.auto_filter.ref = f"A3:P{3 + len(L)}"

    # ======================= CATALOGO =======================
    wc = wb.create_sheet("Catalogo")
    wc["A1"] = "Código"; wc["B1"] = "Nombre del producto"; wc["C1"] = "Fuente del nombre"
    encabezado(wc, 1, ["Código", "Nombre del producto", "Fuente del nombre"])
    for i, cod in enumerate(sorted(set(R.index) | set(maestro.index)), start=2):
        wc.cell(i, 1, cod)
        if cod in maestro.index and maestro[cod]:
            wc.cell(i, 2, maestro[cod]); wc.cell(i, 3, "Maestro de productos Siigo")
        elif cod in nombres.index:
            wc.cell(i, 2, nombres[cod]); wc.cell(i, 3, "Provisional: no está en el maestro; texto de la col. Q sin lote")
        else:
            wc.cell(i, 3, "Pendiente: no está en el maestro ni hay texto en col. Q")
        wc.cell(i, 2).fill = AMAR
        for c in range(1, 4): wc.cell(i, c).font = F(); wc.cell(i, c).border = BORDE
    anchos(wc, [11, 60, 42]); wc.freeze_panes = "A2"

    # ======================= DATOS =======================
    wd = wb.create_sheet("Datos")
    dcols = ["Mes", "Tipo transacción", "Número comprobante", "Fecha elaboración", "Tercero", "Código",
             "Registro Siigo col. Q (original)", "Cantidad", "Lote leído", "FV leída", "Invima leído", "Nota de lectura"]
    encabezado(wd, 1, dcols)
    dd = d.sort_values(["Fecha", "Número comprobante"])
    for i, (_, x) in enumerate(dd.iterrows(), start=2):
        vals = [x["Mes"], x["Tipo transacción"], x["Número comprobante"], x["Fecha"].to_pydatetime(), x["Nombre tercero"],
                x["Código"], x["Nombre"] if isinstance(x["Nombre"], str) else "", float(x["Cantidad"]),
                x["Lote"], x["FV leída"], x["Invima leído"], x["Nota lectura"]]
        for c, v in enumerate(vals, 1):
            cell = wd.cell(i, c, v); cell.font = F(size=9)
        wd.cell(i, 4).number_format = "DD/MM/YYYY"
    anchos(wd, [11, 17, 13, 12, 30, 10, 45, 9, 16, 10, 20, 30])
    wd.freeze_panes = "A2"; wd.auto_filter.ref = f"A1:L{len(dd) + 1}"

    # ======================= RESUMEN =======================
    stats = construir_resumen(d, L, R, fis, no_fis, total_filas, files, meses, nombres)
    stats.update(maestro_n=len(maestro), maestro_path=re.sub(r"^[0-9a-f]{8}-", "", os.path.basename(maestro_path or "")), sin_mov=len(sin_mov),
                 mov_en_maestro=int(fis["En maestro"].sum()), mov_fuera=list(fis.index[~fis["En maestro"]]),
                 prov_fuera=int(sum(c in nombres.index for c in fis.index[~fis["En maestro"]])),
                 filas_plantilla=len(filas_inv))
    wsr = wb.create_sheet("Resumen", 0)
    escribir_resumen(wsr, stats)
    wb.active = 1
    wb.save(salida)
    return stats

def construir_resumen(d, L, R, fis, no_fis, total_filas, files, meses, nombres):
    comp = d["Es compra"]
    q = d["Nombre"].fillna("").str.upper()
    con_lote = d["Lote"] != ""
    est = {
        "archivos": [os.path.basename(f) for f in files], "meses": meses,
        "filas_total": total_filas, "filas_prod": len(d),
        "compras": int(comp.sum()), "ventas": int((~comp).sum()),
        "u_ent": d.loc[comp, "Cantidad"].sum(), "u_sal": d.loc[~comp, "Cantidad"].sum(),
        "codigos": len(R), "fisicos": len(fis), "no_fisicos": list(no_fis.index),
        "con_lote": int(con_lote.sum()), "sin_lote": int((~con_lote).sum()),
        "escrituras": int(d.loc[con_lote, "Nombre"].nunique()), "lotes": len(L),
        "prod_2": int((L[L["Código"].isin(fis.index)].groupby("Código").size() >= 2).sum()),
        "prod_1": int((L[L["Código"].isin(fis.index)].groupby("Código").size() == 1).sum()),
        "prod_0": int(len(fis) - fis.index.isin(L["Código"]).sum()),
        "vencidos_sel": int(L[(L["Orden"] <= 2) & L["Vencido"] & L["Código"].isin(fis.index)].shape[0]),
        "solo_ventas_sel": int(L[(L["Orden"] <= 2) & L["Última compra"].isna() & L["Código"].isin(fis.index)].shape[0]),
        "conf_fv": int(L["Notas"].str.contains("FV distintas").sum()),
        "conf_inv": int(L["Notas"].str.contains("Invima distintos").sum()),
        "inv_compl": int(L["Notas"].str.contains("Invima completado").sum()),
        "sin_fv": int((L["FV"] == "").sum()),
        "varios": int(d["Nota lectura"].str.contains("varios lotes").sum()),
        "fv_inval": int(d["Nota lectura"].str.contains("inválida").sum()),
        "nombres_prov": len(nombres),
        "fmt": [
            ("Lote con 'L:' (dos puntos)", int(q[con_lote].str.contains(r"(^|\s)L\s*:").sum())),
            ("Lote con 'L ' (espacio)", int(q[con_lote].str.contains(r"(^|\n)L\s+[A-Z0-9]").sum())),
            ("Vencimiento con 'FV'", int(q[con_lote].str.contains("FV").sum())),
            ("Vencimiento con 'V' (sin F)", int(q[con_lote].str.contains(r"(^|\n|\s)V\s*\d").sum())),
            ("Año de vencimiento con 2 dígitos (03/28)", int(q[con_lote].str.contains(r"(?:FV|\bV)\s*:?\s*\d{1,2}/\d{2}(?!\d)").sum())),
            ("Año de vencimiento con 4 dígitos (03/2028)", int(q[con_lote].str.contains(r"\d{1,2}/\d{4}").sum())),
            ("Invima con 'RI'", int(q[con_lote].str.contains(r"RI\b|RI:").sum())),
            ("Invima con 'INV'", int(q[con_lote].str.contains("INV").sum())),
            ("Invima sin guion (2016DM0015377)", int(q[con_lote].str.contains(r"\d{4}(?:DM|RD|M)\d").sum())),
            ("Todo en una línea", int((~d.loc[con_lote, "Nombre"].str.contains("\n")).sum())),
            ("En varias líneas (saltos de línea)", int(d.loc[con_lote, "Nombre"].str.contains("\n").sum())),
        ],
        "top": R[R["Físico"]].head(15),
    }
    return est

def escribir_resumen(ws, s):
    ws.column_dimensions["A"].width = 4; ws.column_dimensions["B"].width = 58; ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 16; ws.column_dimensions["E"].width = 16; ws.column_dimensions["F"].width = 16
    r = 1
    def titulo(t, size=12):
        nonlocal r
        c = ws.cell(r, 1, t); c.font = F(bold=True, size=size, color="FFFFFF"); c.fill = GRIS
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6); r += 1
    def texto(t, bold=False):
        nonlocal r
        c = ws.cell(r, 2, t); c.font = F(bold=bold); c.alignment = IZQ
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = max(15, 13 * (len(t) // 105 + 1)); r += 1
    def fila(k, v, fmt=None):
        nonlocal r
        ws.cell(r, 2, k).font = F(); c = ws.cell(r, 3, v); c.font = F(bold=True)
        c.alignment = Alignment(horizontal="right")
        if fmt: c.number_format = fmt
        for col in (2, 3): ws.cell(r, col).border = BORDE
        r += 1
    def blanco():
        nonlocal r; r += 1

    titulo("RESUMEN DEL PROCESAMIENTO – INVENTARIO REAL MEDITIENDA", 14); blanco()
    titulo("1. Archivos procesados")
    texto(f"Informe Siigo 'Movimiento ventas y compras' de DISTRIBUCIONES MEDITIENDA S.A.S, {len(s['archivos'])} archivos: "
          + ", ".join(m.title() for m in s["meses"]) + " de 2026.")
    fila("Filas leídas en total (todas las hojas)", s["filas_total"], "#,##0")
    fila("Filas usadas (Tipo clasificación = Producto)", s["filas_prod"], "#,##0")
    fila("   · Factura de compra (ingresos)", s["compras"], "#,##0")
    fila("   · Factura de venta (salidas)", s["ventas"], "#,##0")
    fila("Unidades ingresadas (suma col. V en compras)", s["u_ent"], "#,##0.##")
    fila("Unidades salidas (suma col. V en ventas)", s["u_sal"], "#,##0.##")
    texto("Se descartaron: filas de 'Formas de pago' y 'Anticipo' (no tienen producto), 'Documento soporte' y líneas de "
          "'Gasto' (honorarios, plataformas, etc.), y las filas de pie de página 'Procesado en…'.")
    blanco()
    titulo("2. Regla 1 – Cálculo de entradas y salidas por código")
    texto("Por cada código (col. P) se sumó la cantidad (col. V) según el tipo de transacción (col. A): "
          "'Factura de compra' = entrada, 'Factura de venta' = salida. La hoja 'Rotacion' tiene el cálculo con fórmulas "
          "SUMIFS sobre la hoja 'Datos', así que se puede auditar fila por fila.")
    texto("Importante: el neto (entradas – salidas) es solo el movimiento de junio a octubre; no incluye el saldo que había "
          "antes de junio. Por eso puede salir negativo y NO reemplaza el conteo físico.", bold=True)
    fila("Códigos distintos con movimiento", s["codigos"], "#,##0")
    fila("Códigos en la plantilla (productos físicos)", s["fisicos"], "#,##0")
    texto("Fuera de la plantilla, por ser servicios o fletes y no productos de bodega: " + ", ".join(s["no_fisicos"]) + ".")
    blanco()
    titulo("3. Regla 2 – Nombre del producto")
    texto(f"Los nombres salen del maestro '{s['maestro_path']}' ({s['maestro_n']:,} códigos), cargado en la hoja 'Catalogo'. "
          "La plantilla busca el nombre por código con una fórmula (INDEX/MATCH): si se corrige un nombre en 'Catalogo', "
          "se actualiza solo en 'Inventario' y en 'Rotacion'.")
    fila("Productos con movimiento que están en el maestro", s["mov_en_maestro"], "#,##0")
    fila("Productos con movimiento que NO están en el maestro", len(s["mov_fuera"]), "#,##0")
    fila("   · de ellos, con nombre provisional (texto de col. Q)", s["prov_fuera"], "#,##0")
    fila("Productos del maestro sin movimiento jun–oct (al final)", s["sin_mov"], "#,##0")
    fila("Total de productos en la plantilla", s["filas_plantilla"], "#,##0")
    texto("Los códigos con movimiento que no están en el maestro se dejaron en la plantilla, en su puesto de rotación, con la "
          "observación 'Código no está en el maestro'. Son 1.597 líneas de factura y varios de los productos que más rotan "
          "(p.ej. DMJ002, DMF003, LBA014). Pueden ser códigos antiguos o duplicados de otro código del maestro "
          "(p.ej. LBA014 y DMA001 tienen el mismo nombre). Conviene revisarlos con Siigo antes del conteo.", bold=True)
    texto("Códigos con movimiento fuera del maestro: " + ", ".join(s["mov_fuera"]) + ".")
    blanco()
    titulo("4. Regla 3 – Homologación de lote, vencimiento e Invima (col. Q)")
    texto(f"De {s['filas_prod']:,} registros de producto, {s['con_lote']:,} traen lote y {s['sin_lote']:,} no (vienen vacíos, "
          f"con 'L N/A' o con solo el nombre o el Invima). Se encontraron {s['escrituras']:,} textos distintos, que quedaron "
          f"en {s['lotes']:,} lotes únicos (código + lote).")
    texto("Formas de escritura encontradas (un mismo registro puede tener varias):", bold=True)
    for k, v in s["fmt"]: fila("   · " + k, v, "#,##0")
    texto("Reglas de homologación aplicadas:", bold=True)
    for t in [
        "Lote: se toma lo que sigue a 'L', 'L:', 'L.' o 'LOTE'. Para unir escrituras se ignoran mayúsculas/minúsculas, guiones, espacios y ceros a la izquierda (0000668290 = 668290).",
        "Vencimiento: todas las formas se llevan a MM/AAAA. Se reconocen 'FV' o 'V' con 03/28, 03/2028, 0328, 032028, 2028/03, 24/03/2028, 20280324 y SEP/2027. Fechas 12/2099 o 12/2080 se tratan como N/A.",
        "Invima: se reconoce 'RI', 'INV', 'INVIMA' o 'RS', o el número suelto aunque no tenga palabra clave. Se escribe AAAA+letras-número-Rn (2016DM0015377 → 2016DM-0015377; .R1 → -R1). 'NO REQUIERE' y 'N/A' se conservan.",
        "Si un mismo lote aparece con distinta FV o distinto Invima en diferentes facturas, se usa el valor más repetido y se deja la nota en la hoja 'Lotes'.",
        "Si a un lote le falta el Invima, se completa con el Invima más usado por ese mismo código (queda anotado).",
        "Si un registro trae varios lotes en el mismo texto, se toma el primero y se marca para verificar.",
    ]: texto("• " + t)
    fila("Lotes con FV distinta entre facturas (resuelto por mayoría)", s["conf_fv"], "#,##0")
    fila("Lotes con Invima distinto entre facturas (resuelto por mayoría)", s["conf_inv"], "#,##0")
    fila("Lotes con Invima completado desde el producto", s["inv_compl"], "#,##0")
    fila("Lotes sin fecha de vencimiento legible", s["sin_fv"], "#,##0")
    fila("Registros con FV imposible en origen (p.ej. 80/2027, 05/202)", s["fv_inval"], "#,##0")
    fila("Registros con varios lotes en un solo texto", s["varios"], "#,##0")
    blanco()
    titulo("5. Regla 4 – Selección de los 2 lotes por producto")
    texto("Para cada código se ordenaron sus lotes así: primero los que tienen Factura de compra, del ingreso más reciente al "
          "más antiguo (fecha de elaboración); si dos lotes ingresaron el mismo día, va primero el de vencimiento más lejano. "
          "Después van los lotes que solo aparecen en ventas (pudieron entrar antes de junio), del más reciente al más antiguo. "
          "Los dos primeros son la Opción 1 y la Opción 2. Si hay un solo lote, la Opción 2 queda en blanco.")
    fila("Productos con 2 lotes", s["prod_2"], "#,##0")
    fila("Productos con 1 lote (Opción 2 en blanco)", s["prod_1"], "#,##0")
    fila("Productos sin lote en los informes", s["prod_0"], "#,##0")
    fila("Lotes seleccionados que ya están vencidos (marcados en rojo)", s["vencidos_sel"], "#,##0")
    fila("Lotes seleccionados que solo aparecen en ventas", s["solo_ventas_sel"], "#,##0")
    blanco()
    titulo("6. Regla 5 – Orden por rotación")
    texto("La prioridad es el ranking por total movido (unidades ingresadas + unidades salidas, col. V). Si hay empate, va primero "
          "el producto con más facturas. El producto 1 es el de mayor rotación y queda arriba en la hoja 'Inventario'. "
          "Nota: las unidades dependen de cómo se factura cada producto (unidad, caja, paquete).")
    texto("Los 15 productos con mayor rotación:", bold=True)
    hdrs = [("A", "#"), ("B", "Código"), ("C", "Entradas"), ("D", "Salidas"), ("E", "Total movido"), ("F", "N.º facturas")]
    for col, h in hdrs:
        c = ws[f"{col}{r}"]; c.value = h; c.font = F(bold=True, color="FFFFFF"); c.fill = TEAL; c.border = BORDE
    r += 1
    for cod, x in s["top"].iterrows():
        for col, v in zip("ABCDEF", [int(x["Ranking"]), cod, x["Entradas"], x["Salidas"], x["Total movido"], int(x["Mov_total"])]):
            c = ws[f"{col}{r}"]; c.value = v; c.font = F(); c.border = BORDE
            if col in "CDE": c.number_format = "#,##0.##"
        r += 1
    blanco()
    titulo("7. Cómo usar la plantilla")
    for t in [
        "Hoja 'Inventario': cuente en bodega y escriba en 'Cantidad por lote' (celdas amarillas) las unidades de cada lote. 'Total contado' suma las dos opciones del producto.",
        "Si encuentra un lote distinto a los dos sugeridos, sobrescriba Lote / FV / Invima en esa fila o agregue una nota en 'Observación'.",
        "Hoja 'Rotacion': entradas, salidas y número de facturas por código (con fórmulas).",
        "Hoja 'Lotes': todos los lotes de cada producto, ya homologados, con el criterio de selección y las notas de lo que se corrigió.",
        "Hoja 'Catalogo': maestro de códigos y nombres; corrija aquí un nombre y se actualiza en todo el libro.",
        "Hoja 'Datos': las líneas de producto de los 5 informes, con el lote, la FV y el Invima leídos de cada una.",
    ]: texto("• " + t)

if __name__ == "__main__":
    # uso: python generar_excel.py salida.xlsx [--maestro CODIGOS.xlsx] informe1.xlsx informe2.xlsx ...
    args = sys.argv[1:]
    salida = args.pop(0); maestro = None
    if "--maestro" in args:
        i = args.index("--maestro"); maestro = args[i + 1]; del args[i:i + 2]
    st = main(args, salida, maestro)
    print({k: v for k, v in st.items() if k not in ("top", "fmt", "mov_fuera")})
