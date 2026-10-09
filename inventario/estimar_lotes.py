"""Estima la cantidad por lote sobre la base de inventario (2 lotes por código) con los informes Siigo jun–oct.

Estimado = comprado del lote − vendido del lote − ventas sin lote asignadas (FIFO), mínimo 0.
uso: python estimar_lotes.py BASE.xlsx SALIDA.xlsx informe1.xlsx informe2.xlsx ...
"""
import sys
import pandas as pd
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from procesar_inventario import cargar, analizar, clave_lote, HOY

ARIAL = "Arial"
VERDE = PatternFill("solid", fgColor="2E7D32")
fino = Side(style="thin", color="7F7F7F")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
IZQ = Alignment(horizontal="left", vertical="center", wrap_text=True)

def calcular(d):
    """Por (código, clave lote): comprado, vendido y ventas sin lote asignadas por FIFO."""
    cl = d[d["Clave lote"] != ""]
    comp = cl[cl["Es compra"]].groupby(["Código", "Clave lote"]).agg(comp=("Cantidad", "sum"), f_ini=("Fecha", "min"))
    vend = cl[~cl["Es compra"]].groupby(["Código", "Clave lote"])["Cantidad"].sum().rename("vend")
    t = comp.join(vend, how="outer").fillna({"comp": 0, "vend": 0})
    t["sin_lote"] = 0.0
    sin_lote = d[(d["Clave lote"] == "") & ~d["Es compra"]].groupby("Código")["Cantidad"].sum()
    sobrante = {}
    for cod, pend in sin_lote.items():
        if cod not in t.index.get_level_values(0): sobrante[cod] = pend; continue
        lotes = t.loc[cod]
        lotes = lotes[lotes["comp"] > 0].sort_values("f_ini")  # primero en entrar, primero en salir
        for k, x in lotes.iterrows():
            if pend <= 0: break
            disp = max(0.0, x["comp"] - x["vend"])
            usa = min(disp, pend)
            t.loc[(cod, k), "sin_lote"] = usa; pend -= usa
        if pend > 0: sobrante[cod] = pend
    return t, sin_lote, sobrante

def main(base, salida, files):
    d, _ = cargar(files)
    d, *_ = analizar(d)
    t, sin_lote, sobrante = calcular(d)

    wb = load_workbook(base)
    ws = wb.worksheets[0]
    hdr_ref = ws["H2"]; g_ref = ws["E1"]
    ws.merge_cells("I1:N1"); ws["I1"] = "ESTIMACIÓN SEGÚN INFORMES JUN–OCT 2026"
    heads = ["Comprado del lote", "Vendido con este lote", "Ventas sin lote asignadas (FIFO)",
             "Cantidad estimada", "Confianza", "Diferencia (conteo − estimado)"]
    for c in range(9, 15):
        cell = ws.cell(1, c); cell.fill = VERDE; cell.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=g_ref.font.sz or 11)
        cell.alignment = CENTRO; cell.border = BORDE
        h = ws.cell(2, c, heads[c - 9]); h.font = copy(hdr_ref.font); h.fill = VERDE; h.alignment = CENTRO; h.border = BORDE
    for col, w in zip("IJKLMN", [11, 11, 13, 11, 34, 13]): ws.column_dimensions[col].width = w
    ws.row_dimensions[2].height = 45

    cod = None; stats = {"alta": 0, "media": 0, "sin": 0, "vacio": 0}
    ultima = ws.max_row
    for r in range(3, ultima + 1):
        v = ws.cell(r, 2).value
        if v: cod = str(v).strip().upper()
        lote = ws.cell(r, 5).value
        fv = str(ws.cell(r, 6).value or "")
        for c in range(9, 15):
            cell = ws.cell(r, c); cell.border = BORDE; cell.font = Font(name=ARIAL, size=10)
            cell.alignment = IZQ if c == 13 else CENTRO
        ws.cell(r, 14, f'=IF(OR(H{r}="",L{r}=""),"",H{r}-L{r})')
        if not lote: stats["vacio"] += 1; continue
        k = clave_lote(str(lote).strip().upper())
        x = t.loc[(cod, k)] if (cod, k) in t.index else None
        comp = float(x["comp"]) if x is not None else 0.0
        vend = float(x["vend"]) if x is not None else 0.0
        asig = float(x["sin_lote"]) if x is not None else 0.0
        ws.cell(r, 9, comp); ws.cell(r, 10, vend); ws.cell(r, 11, asig)
        ws.cell(r, 12, f'=IF(I{r}=0,"",MAX(0,I{r}-J{r}-K{r}))')
        vencido = len(fv) == 7 and fv[2] == "/" and fv.replace("/", "").isdigit() and \
            pd.Timestamp(int(fv[3:]), int(fv[:2]), 1) + pd.offsets.MonthEnd(0) < HOY
        if comp == 0:
            conf = "Sin base: lote comprado antes de junio" if vend > 0 else "Sin base: lote sin movimiento en los informes"
            stats["sin"] += 1
        elif vend > comp:
            conf = "Media: se vendió más de lo comprado en el periodo (había saldo antes de junio)"; stats["media"] += 1
        elif asig > 0:
            conf = "Media: incluye ventas sin lote asignadas por FIFO"; stats["media"] += 1
        else:
            conf = "Alta: compras y ventas del lote identificadas"; stats["alta"] += 1
        if vencido: conf += " · LOTE VENCIDO"
        ws.cell(r, 13, conf)
        for c in (9, 10, 11, 12, 14): ws.cell(r, c).number_format = "#,##0.##;-#,##0.##;0"
    ws.conditional_formatting.add(f"N3:N{ultima}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

    # ---------- hoja explicativa
    we = wb.create_sheet("Metodo estimacion")
    we.column_dimensions["A"].width = 110
    lineas = [
        ("CÓMO SE CALCULÓ LA CANTIDAD ESTIMADA POR LOTE", True),
        ("", False),
        ("Fuente: informes Siigo 'Movimiento ventas y compras' de junio a octubre de 2026 (solo líneas de producto). Columna A: "
         "Factura de compra = entrada; Factura de venta = salida. Columna V = cantidad. Columna Q = lote (homologado como en la plantilla).", False),
        ("", False),
        ("1. Comprado del lote: suma de unidades en facturas de compra con ese código y ese lote.", False),
        ("2. Vendido con este lote: suma de unidades en facturas de venta con ese código y ese lote.", False),
        ("3. Ventas sin lote asignadas (FIFO): ventas del producto que no traen lote en la col. Q. Se descuentan primero del lote "
         "que entró antes (primero en entrar, primero en salir) y sin pasar de lo que le quedaba a cada lote.", False),
        ("4. Cantidad estimada = Comprado − Vendido − Ventas sin lote asignadas, con un mínimo de 0. Es una fórmula, así que se puede ajustar.", False),
        ("5. Diferencia = Cantidad por lote (conteo físico) − Cantidad estimada. Si sale negativa (en rojo), en bodega hay menos de lo esperado.", False),
        ("", False),
        ("Nivel de confianza", True),
        (f"• Alta ({stats['alta']} lotes): el lote se compró en el periodo y todas sus ventas tienen lote. El estimado es el más confiable.", False),
        (f"• Media ({stats['media']} lotes): el lote se compró en el periodo, pero se le asignaron ventas sin lote, o se vendió más de lo comprado "
         "(lo que indica que ya había unidades de ese lote antes de junio y el saldo real puede ser mayor).", False),
        (f"• Sin base ({stats['sin']} lotes): el lote no tiene compras entre junio y octubre (entró antes de junio). Sin el saldo al 1 de junio no "
         "se puede estimar; queda en blanco y se debe contar.", False),
        (f"• Filas sin lote en la base: {stats['vacio']}.", False),
        ("", False),
        ("Limitaciones", True),
        ("• No hay saldo inicial al 1 de junio, así que el estimado solo cubre lo que entró en el periodo.", False),
        ("• Ajustes de inventario, devoluciones, remisiones y traslados no están en estos informes.", False),
        ("• Si en una venta se escribió un lote equivocado, la venta queda en ese lote y no en el real.", False),
        (f"• Ventas sin lote que no se pudieron asignar porque no había saldo comprado en el periodo: "
         f"{sum(sobrante.values()):,.0f} unidades en {len(sobrante)} productos (seguramente salieron de stock anterior a junio).", False),
        ("• Es una cantidad aproximada para orientar el conteo; el dato válido es el conteo físico de la columna 'Cantidad por lote'.", True),
    ]
    for i, (txt, b) in enumerate(lineas, 1):
        c = we.cell(i, 1, txt); c.font = Font(name=ARIAL, bold=b, size=12 if i == 1 else 10); c.alignment = IZQ
    wb.save(salida)
    return stats, sobrante

if __name__ == "__main__":
    st, sob = main(sys.argv[1], sys.argv[2], sys.argv[3:])
    print(st, len(sob), sum(sob.values()))
